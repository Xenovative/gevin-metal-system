"""取消訂單：從主資料庫刪除單據，存檔至 cancelled_orders.xlsx，單號永不重用。

取代了舊版「作廢」（保留單據＋沖銷記錄）。取消流程：
1. 先將完整單據快照附加到 output/reports/cancelled_orders.xlsx（寫入失敗則中止，資料庫不變）
2. 刪除 invoices / invoice_line_items / inventory_movements / cash_movements 中該單號所有記錄
3. 寫入 cancelled_invoices 封鎖表，該單號之後不可再被系統分配
"""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook, Workbook

from auth import log_audit, require_admin
from config import REPORT_DIR

logger = logging.getLogger(__name__)

ARCHIVE_PATH = REPORT_DIR / "cancelled_orders.xlsx"
ARCHIVE_SHEET = "取消訂單"
ARCHIVE_HEADERS = [
    "單號", "交易性質", "日期", "客戶", "客戶電話", "區塊",
    "貨品", "成色", "重量(克)", "重量(両)", "重量(安士)", "單價", "金額",
    "貨幣", "倉存存取", "倉存位置", "經手人", "原備註", "取消時間", "取消操作者",
]


def _append_cancelled_archive(invoice, line_items, cancelled_by, cancelled_at):
    """Append one row per line item to cancelled_orders.xlsx (created if missing)."""
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    if ARCHIVE_PATH.exists():
        wb = load_workbook(ARCHIVE_PATH)
        ws = wb[ARCHIVE_SHEET] if ARCHIVE_SHEET in wb.sheetnames else wb.active
        ws.title = ARCHIVE_SHEET
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = ARCHIVE_SHEET
        ws.append(ARCHIVE_HEADERS)

    base = [
        invoice.invoice_no,
        invoice.transaction_type or "",
        invoice.transaction_date.strftime("%Y-%m-%d") if invoice.transaction_date else "",
        invoice.customer_name or "",
        invoice.customer_phone or "",
    ]
    tail = [
        invoice.invoice_currency or "",
        invoice.source_location or "",
        invoice.destination_location or "",
        invoice.handler or "",
        invoice.notes or "",
        cancelled_at.strftime("%Y-%m-%d %H:%M"),
        cancelled_by,
    ]
    rows = []
    items = list(line_items)
    if not items:
        items = [None]
    for item in items:
        if item is None:
            rows.append(base + ["", "", "", "", "", "", "", ""] + tail)
            continue
        rows.append(base + [
            "對換" if (item.section or "") == "exchange" else "主項",
            item.item_type or "",
            item.quality or "",
            item.weight_gram if item.weight_gram is not None else "",
            item.weight_tael if item.weight_tael is not None else "",
            item.weight_oz if getattr(item, "weight_oz", None) is not None else "",
            item.unit_price if item.unit_price is not None else "",
            item.amount if item.amount is not None else "",
        ] + tail)
    for row in rows:
        ws.append(row)
    wb.save(ARCHIVE_PATH)
    return len(rows)


def cancelled_archive_file():
    """Return the archive Excel path if it exists, else None."""
    return str(ARCHIVE_PATH) if ARCHIVE_PATH.exists() else None


def load_cancelled_archive_rows():
    """Read every data row from cancelled_orders.xlsx as a list of dicts."""
    if not ARCHIVE_PATH.exists():
        return []
    wb = load_workbook(ARCHIVE_PATH, data_only=True)
    try:
        ws = wb[ARCHIVE_SHEET] if ARCHIVE_SHEET in wb.sheetnames else wb.active
        rows = []
        for i, raw in enumerate(ws.iter_rows(values_only=True)):
            if not raw or raw[0] is None:
                continue
            if i == 0 and str(raw[0]).strip() == ARCHIVE_HEADERS[0]:
                continue
            rec = {}
            for idx, header in enumerate(ARCHIVE_HEADERS):
                value = raw[idx] if idx < len(raw) else None
                rec[header] = "" if value is None else value
            rows.append(rec)
        return rows
    finally:
        wb.close()


def search_cancelled_archive(query=""):
    """Filter archive rows by serial, customer, goods, handler, etc.

    Empty query returns the full list.
    """
    rows = load_cancelled_archive_rows()
    needle = str(query or "").strip().lower()
    if not needle:
        return rows
    search_keys = (
        "單號", "交易性質", "客戶", "客戶電話", "貨品", "成色",
        "經手人", "取消操作者", "原備註",
    )
    matched = []
    for rec in rows:
        blob = " ".join(str(rec.get(k, "") or "") for k in search_keys).lower()
        if needle in blob:
            matched.append(rec)
    return matched


def cancel_invoice(session, invoice_no, admin_user):
    from database import (
        CancelledInvoice,
        CashMovement,
        InventoryMovement,
        Invoice,
        InvoiceLineItem,
    )

    err = require_admin(admin_user)
    if err:
        return err

    invoice_no = (invoice_no or "").strip()
    if not invoice_no:
        return "❌ 請填寫單號"

    already = (
        session.query(CancelledInvoice)
        .filter(CancelledInvoice.invoice_no == invoice_no)
        .first()
    )
    if already:
        return f"❌ 單號「{invoice_no}」已於 {already.cancelled_at:%Y-%m-%d %H:%M} 取消，不可重複取消"

    invoice = session.query(Invoice).filter(Invoice.invoice_no == invoice_no).first()
    if not invoice:
        return f"❌ 找不到單號「{invoice_no}」"

    line_items = (
        session.query(InvoiceLineItem)
        .filter(InvoiceLineItem.invoice_id == invoice.id)
        .order_by(InvoiceLineItem.sort_order, InvoiceLineItem.id)
        .all()
    )
    # Expunge snapshot data before deletes so rows stay readable afterwards
    for obj in [invoice, *line_items]:
        session.expunge(obj)

    cancelled_by = admin_user.get("display_name", "")
    cancelled_at = datetime.now()

    # Archive first — if this fails the database stays untouched
    try:
        archived_rows = _append_cancelled_archive(
            invoice, line_items, cancelled_by, cancelled_at
        )
    except Exception as exc:
        session.rollback()
        logger.exception("Cancel archive failed for %s", invoice_no)
        return f"❌ 存檔 cancelled_orders.xlsx 失敗，未刪除任何記錄：{exc}"

    cash_deleted = (
        session.query(CashMovement)
        .filter(CashMovement.invoice_no == invoice_no)
        .delete(synchronize_session=False)
    )
    inv_deleted = (
        session.query(InventoryMovement)
        .filter(InventoryMovement.invoice_no == invoice_no)
        .delete(synchronize_session=False)
    )
    items_deleted = (
        session.query(InvoiceLineItem)
        .filter(InvoiceLineItem.invoice_id == invoice.id)
        .delete(synchronize_session=False)
    )
    session.query(Invoice).filter(Invoice.id == invoice.id).delete(
        synchronize_session=False
    )

    session.add(
        CancelledInvoice(
            invoice_no=invoice.invoice_no,
            transaction_type=invoice.transaction_type,
            customer_name=invoice.customer_name,
            transaction_date=invoice.transaction_date,
            total_amount=invoice.total_amount,
            invoice_currency=invoice.invoice_currency,
            cancelled_by=cancelled_by,
            cancelled_at=cancelled_at,
        )
    )
    log_audit(session, admin_user, "cancel_invoice", "invoice", invoice_no, {
        "line_items_deleted": items_deleted,
        "inventory_movements_deleted": inv_deleted,
        "cash_movements_deleted": cash_deleted,
        "archived_rows": archived_rows,
    })
    session.commit()
    logger.info(
        "Cancelled %s: invoice+items/movements deleted, serial blocked", invoice_no
    )
    return (
        f"✅ 已取消單號 {invoice_no}\n"
        f"已從資料庫刪除：貨品 {items_deleted} 項、倉存記錄 {inv_deleted} 筆、"
        f"現金記錄 {cash_deleted} 筆\n"
        f"已存檔至 {Path(ARCHIVE_PATH).name}（{archived_rows} 列）\n"
        f"此單號永不重用"
    )
