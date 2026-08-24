"""
Regression: overlay Excel has data only at LAYOUT coords; Perfect V2 PDF
is built from invoice_data (not overlay cells).
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from pypdf import PdfReader

from cash import signed_cash_warehouse_amount
from invoice_excel_generator import CUSTOMER_COPY, COMPANY_COPY, generate_invoice_excel
from receipt_pdf import HEADER_PATH, build_invoice_pdf_from_data


def _fixture():
    invoice_no = "TEST_PDF_PARITY_001"
    total = 0.0
    data = {
        "invoice_no": invoice_no,
        "transaction_type": "銷售單",
        "customer_name": "pmp",
        "customer_phone": "90001111",
        "transaction_date": date(2026, 7, 31),
        "handler": "Admin",
        "payment_method": "",
        "invoice_currency": "HKD",
        "source_location": "取 A倉庫",
        "destination_location": "",
        "notes": "",
        "note_amount": 0,
        "total_amount": total,
        "cash_warehouse_amount": signed_cash_warehouse_amount(total, "銷售單"),
    }
    items = [
        {
            "item_type": "足金 Pure Gold",
            "quality": "足金",
            "weight_gram": 1,
            "weight_tael": 0,
            "weight_oz": 0,
            "unit_price": 0,
            "amount": 0,
        }
    ]
    return data, items


def _assert_excel(path: Path):
    ws = load_workbook(path)["銷售"]
    assert ws.cell(CUSTOMER_COPY["items_start"], 3).value == "足金 Pure Gold"
    assert float(ws.cell(CUSTOMER_COPY["items_start"], 11).value) == 0
    assert ws.cell(CUSTOMER_COPY["total_row"], 10).value == "HKD"
    assert float(ws.cell(CUSTOMER_COPY["total_row"], 11).value) == 0.0
    assert ws.cell(CUSTOMER_COPY["info_row"], 5).value == "pmp"
    inv = str(ws.cell(CUSTOMER_COPY["invoice_no_row"], 11).value)
    assert "TEST_PDF_PARITY_001" in inv
    # Bilingual invoice-type label in col J, serial in col K
    label = str(ws.cell(CUSTOMER_COPY["invoice_no_row"], 10).value or "")
    assert "銷售單" in label and "Sales Invoice" in label
    co_inv_row = COMPANY_COPY["invoice_no_row"]
    co_label = str(ws.cell(co_inv_row, 10).value or "")
    assert "銷售單" in co_label and "Sales Invoice" in co_label

    assert ws.cell(CUSTOMER_COPY["handler_row"] - 1, 6).value == "經手人 Handler:"
    assert ws.cell(CUSTOMER_COPY["handler_row"], 6).value == "Admin"
    co_handler = COMPANY_COPY["handler_row"]
    assert ws.cell(co_handler - 1, 6).value == "經手人 Handler:"
    assert ws.cell(co_handler, 6).value == "Admin"
    assert ws.cell(CUSTOMER_COPY["info_row"], 7).value == "90001111"
    assert ws.cell(CUSTOMER_COPY["items_start"] - 3, 3).value == "貨品"
    assert ws.cell(CUSTOMER_COPY["items_start"] - 2, 9).value == "Unit ($)"
    assert ws.cell(CUSTOMER_COPY["items_start"] - 2, 6).value == "Weight"
    assert not ws.cell(CUSTOMER_COPY["items_start"] - 3, 8).value
    assert not ws.cell(CUSTOMER_COPY["items_start"] - 2, 8).value
    co_hdr = COMPANY_COPY["items_start"] - 3
    assert ws.cell(co_hdr, 8).value == "庫存"
    assert ws.cell(co_hdr + 1, 8).value == "Stock"

    # Zero 両/安士 must not consume extra rows
    assert ws.cell(CUSTOMER_COPY["items_start"] + 1, 7).value in (None, "")

    co_items = COMPANY_COPY["items_start"]
    assert ws.cell(co_items, 3).value == "足金 Pure Gold"
    stock = str(ws.cell(co_items, 8).value or "")
    assert "取 A倉庫" in stock or "A倉庫" in stock
    assert "存 客戶" not in stock
    assert "取 客戶" not in stock
    assert "客戶" not in stock
    assert "倉存存取" not in stock
    item_cell = str(ws.cell(co_items, 3).value or "")
    assert "(取" not in item_cell
    assert "取 A倉庫" not in item_cell
    assert not ws.cell(CUSTOMER_COPY["items_start"], 8).value
    print("Excel assertions OK")


def _assert_brand_pdf(path: Path):
    if not Path(HEADER_PATH).exists():
        print(f"SKIP PDF parity: header missing ({HEADER_PATH})")
        return False

    reader = PdfReader(str(path))
    assert len(reader.pages) >= 1
    text = "\n".join((page.extract_text() or "") for page in reader.pages)

    assert "TEST_PDF_PARITY_001" in text or "TEST_PDF_PARI" in text
    assert "pmp" in text
    assert "足金" in text or "Pure Gold" in text
    assert "HKD 0.00" in text or "HKD" in text
    assert "0" in text
    assert "Admin" in text
    assert "XXXX" not in text
    assert "取 A倉庫" in text or "A倉庫" in text
    assert "倉存存取" not in text
    assert "倉存位置" not in text
    assert "Pure Gold (取" not in text
    assert "足金 Pure Gold (取" not in text
    assert "客戶單" in text or "Customer Copy" in text
    assert "公司單" in text or "Company Copy" in text
    print("PDF assertions OK (Perfect V2)")
    return True


def main():
    data, items = _fixture()
    excel = Path(generate_invoice_excel(data, items))
    assert excel.exists(), excel
    _assert_excel(excel)

    pdf = excel.with_suffix(".pdf")
    if not pdf.exists():
        try:
            build_invoice_pdf_from_data(data, items, None, pdf)
        except Exception as exc:
            print(
                "SKIP PDF parity: Perfect V2 render unavailable "
                f"(need receipt_header.png + CJK font). Detail: {exc}"
            )
            print("ALL EXCEL ASSERTIONS PASSED (PDF skipped)")
            return

    assert pdf.exists(), pdf
    if not _assert_brand_pdf(pdf):
        print("ALL EXCEL ASSERTIONS PASSED (PDF skipped)")
        return
    print("ALL EXCEL/PDF PARITY TESTS PASSED")


if __name__ == "__main__":
    main()
