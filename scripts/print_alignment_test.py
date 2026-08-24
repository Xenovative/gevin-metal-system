"""Print-alignment tests: overlay Excel cells + 100% A4 setup for all tx types."""
from datetime import date
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from cash import signed_cash_warehouse_amount
from invoice_excel_generator import (
    COL_WIDTHS,
    COMPANY_COPY,
    CUSTOMER_COPY,
    PRINT_AREA,
    estimate_items_block_rows,
    fit_items_for_print,
    generate_invoice_excel,
)


FORBIDDEN_LABELS = (
    "XXX",
)


def _base_data(invoice_no, tx_type, total, currency="HKD"):
    return {
        "invoice_no": invoice_no,
        "transaction_type": tx_type,
        "customer_name": "對齊測試客戶",
        "customer_phone": "91234567",
        "transaction_date": date(2026, 8, 11),
        "handler": "倉管",
        "payment_method": (
            '[{"method":"現金","amount":100,"currency":"HKD"},'
            '{"method":"轉帳","amount":50,"currency":"HKD"}]'
        ),
        "invoice_currency": currency,
        "source_location": "取 A倉庫",
        "destination_location": "",
        "notes": "備註測試",
        "note_amount": 0,
        "total_amount": total,
        "cash_warehouse_amount": signed_cash_warehouse_amount(total, tx_type),
    }


def _sample_item(amount=150.0):
    return {
        "item_type": "黃金 Gold",
        "quality": "999.9",
        "weight_gram": 10,
        "weight_tael": 0.267,
        "weight_oz": None,
        "unit_price": 15,
        "amount": amount,
    }


def _assert_a4(ws):
    assert int(ws.page_setup.paperSize or 0) == int(ws.PAPERSIZE_A4)
    assert int(ws.page_setup.scale or 0) == 100
    fit_pr = ws.sheet_properties.pageSetUpPr
    assert fit_pr is None or fit_pr.fitToPage in (False, 0, None)
    area = (ws.print_area or "").replace("$", "")
    assert PRINT_AREA.replace("$", "") in area or "A1:K50" in area
    merged = {str(r).replace("$", "") for r in ws.merged_cells.ranges}
    for ref in ("F8:G8", "F9:G9", "F35:G35", "F36:G36"):
        assert ref not in merged, f"weight header still merged: {ref}"
    for letter, width in COL_WIDTHS.items():
        actual = float(ws.column_dimensions[letter].width or 0)
        assert abs(actual - width) < 0.05, (letter, actual, width)
    assert abs(float(ws.row_dimensions[21].height or 0) - 27.75) < 0.05
    assert abs(float(ws.row_dimensions[45].height or 0) - 27.75) < 0.05
    assert abs(float(ws.row_dimensions[37].height or 0) - 24.75) < 0.05


def _assert_no_static_labels(ws):
    for row in ws.iter_rows(min_row=1, max_row=54, max_col=11):
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value)
            for bad in FORBIDDEN_LABELS:
                assert bad not in text, f"{cell.coordinate}: {text}"


def _assert_copy_cells(ws, layout, invoice_no, customer, signed_total, currency, has_amount=True, handler="倉管"):
    inv = str(ws.cell(row=layout["invoice_no_row"], column=11).value or "")
    assert invoice_no in inv, inv
    assert ws.cell(row=layout["info_row"], column=5).value == customer
    assert ws.cell(row=layout["info_row"], column=3).value == "客戶 Customer :"
    assert str(ws.cell(row=layout["total_row"], column=3).value or "").startswith("付款方式")
    assert abs(float(ws.row_dimensions[layout["total_row"]].height or 0) - 27.75) < 0.05
    assert ws.cell(row=layout["items_start"], column=3).value  # item type
    assert ws.cell(row=layout["handler_row"], column=6).value == handler
    assert ws.cell(row=layout["handler_row"] - 1, column=6).value == "經手人 Handler:"
    label = str(ws.cell(row=layout["invoice_no_row"], column=10).value or "")
    assert "單" in label or "Invoice" in label or "Note" in label, label
    assert ws.cell(row=layout["info_row"], column=7).value == "91234567"
    assert ws.cell(row=layout["items_start"] - 3, column=3).value == "貨品"
    assert ws.cell(row=layout["items_start"] - 2, column=3).value == "Item"
    weight_en = ws.cell(row=layout["items_start"] - 2, column=6).value
    assert weight_en in ("Weight", "Gross"), weight_en
    is_company = layout["items_start"] > CUSTOMER_COPY["items_start"]
    if is_company:
        assert abs(float(ws.row_dimensions[layout["items_start"]].height or 0) - 24.75) < 0.05
    hdr_stock = ws.cell(row=layout["items_start"] - 3, column=8).value
    en_stock = ws.cell(row=layout["items_start"] - 2, column=8).value
    if is_company:
        assert hdr_stock == "庫存", hdr_stock
        assert en_stock == "Stock", en_stock
    else:
        assert hdr_stock in (None, ""), hdr_stock
        assert en_stock in (None, ""), en_stock
    cash_cn = (
        ws.cell(row=layout["items_start"] - 3, column=11).value
        or ws.cell(row=layout["items_start"] - 3, column=10).value
    )
    cash_en = (
        ws.cell(row=layout["items_start"] - 2, column=11).value
        or ws.cell(row=layout["items_start"] - 2, column=10).value
    )
    assert cash_cn in (None, ""), cash_cn
    assert cash_en in (None, ""), cash_en
    if has_amount:
        assert ws.cell(row=layout["total_row"], column=10).value == currency
        assert float(ws.cell(row=layout["total_row"], column=11).value) == float(signed_total)
        assert float(ws.cell(row=layout["items_start"], column=11).value) == float(signed_total)


def _company_layout():
    return dict(COMPANY_COPY)


def test_sales():
    data = _base_data("ALIGN_S1", "銷售單", 150.0)
    items = [_sample_item(150.0)]
    path = generate_invoice_excel(data, items)
    ws = load_workbook(path)["銷售"]
    _assert_a4(ws)
    _assert_no_static_labels(ws)
    signed = signed_cash_warehouse_amount(150.0, "銷售單")
    _assert_copy_cells(ws, CUSTOMER_COPY, "ALIGN_S1", "對齊測試客戶", signed, "HKD")
    _assert_copy_cells(ws, _company_layout(), "ALIGN_S1", "對齊測試客戶", signed, "HKD")
    stock = str(ws.cell(row=_company_layout()["items_start"], column=8).value or "")
    assert "A倉庫" in stock
    assert "存 客戶" not in stock
    assert "取 客戶" not in stock
    assert not ws.cell(row=CUSTOMER_COPY["items_start"], column=8).value
    assert ws.cell(row=CUSTOMER_COPY["payment_row"], column=3).value
    print("OK 銷售", path)


def test_purchase():
    data = _base_data("ALIGN_P1", "購入單", 200.0)
    data["source_location"] = ""
    data["destination_location"] = "存 A倉庫"
    items = [{**_sample_item(200.0), "amount": 200.0, "unit_price": 20}]
    path = generate_invoice_excel(data, items)
    ws = load_workbook(path)["購入單"]
    _assert_a4(ws)
    _assert_no_static_labels(ws)
    signed = signed_cash_warehouse_amount(200.0, "購入單")
    assert signed == -200.0
    _assert_copy_cells(ws, CUSTOMER_COPY, "ALIGN_P1", "對齊測試客戶", signed, "HKD")
    _assert_copy_cells(ws, _company_layout(), "ALIGN_P1", "對齊測試客戶", signed, "HKD")
    co_start = _company_layout()["items_start"]
    stock = " ".join(str(ws.cell(row=co_start + i, column=8).value or "") for i in range(3))
    assert "存 A倉庫" in stock, stock
    assert "取 客戶" not in stock
    print("OK 購入", path)


def test_exchange():
    data = _base_data("ALIGN_T1", "兌料單", 80.0)
    main = [_sample_item(80.0)]
    exchange = [{
        "item_type": "純銀 Silver",
        "quality": "999",
        "weight_gram": 30,
        "weight_tael": None,
        "weight_oz": None,
        "unit_price": None,
        "amount": None,
    }]
    path = generate_invoice_excel(data, main, exchange)
    ws = load_workbook(path)["兌料單"]
    _assert_a4(ws)
    _assert_no_static_labels(ws)
    assert "ALIGN_T1" in str(ws.cell(row=CUSTOMER_COPY["invoice_no_row"], column=11).value)
    notes_val = ws.cell(row=CUSTOMER_COPY["notes_row"], column=5).value
    assert notes_val == "備註測試", notes_val
    print("OK 兌料", path)


def test_delivery():
    data = _base_data("ALIGN_D1", "交收單", 0.0)
    data["payment_method"] = ""
    items = [{
        "item_type": "黃金 Gold",
        "quality": "999.9",
        "weight_gram": 5,
        "weight_tael": None,
        "weight_oz": None,
        "unit_price": None,
        "amount": None,
    }]
    path = generate_invoice_excel(data, items)
    ws = load_workbook(path)["交收單"]
    _assert_a4(ws)
    _assert_no_static_labels(ws)
    assert "ALIGN_D1" in str(ws.cell(row=CUSTOMER_COPY["invoice_no_row"], column=11).value)
    assert ws.cell(row=CUSTOMER_COPY["items_start"], column=3).value == "黃金 Gold"
    assert ws.cell(row=CUSTOMER_COPY["total_row"], column=11).value in (None, "")
    print("OK 交收", path)


def test_overflow_clamp():
    data = _base_data("ALIGN_OV1", "銷售單", 10.0)
    items = []
    for i in range(12):
        items.append({
            "item_type": f"Item{i}",
            "quality": "999",
            "weight_gram": 1,
            "weight_tael": 0.1,
            "weight_oz": 0.03,
            "unit_price": 1,
            "amount": 1,
        })

    fitted, _, warning = fit_items_for_print(items)
    assert warning
    max_rows = CUSTOMER_COPY["notes_row"] - CUSTOMER_COPY["items_start"]
    assert estimate_items_block_rows(fitted) <= max_rows

    path = generate_invoice_excel(data, items)
    ws = load_workbook(path)["銷售"]
    notes_cell = ws.cell(row=CUSTOMER_COPY["notes_row"], column=4).value
    assert notes_cell == "備註測試", notes_cell
    assert data.get("print_warning")
    print("OK overflow", path, "warning=", data["print_warning"])


def main():
    test_sales()
    test_purchase()
    test_exchange()
    test_delivery()
    test_overflow_clamp()
    print("ALL PRINT ALIGNMENT TESTS PASSED")


if __name__ == "__main__":
    main()
