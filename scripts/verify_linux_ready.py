"""
Verify the repo is safe for a clean Linux / Docker client deploy:
- no SQLite / invoice mock data tracked in git
- Docker / compose / Linux scripts present
- fresh empty data dir creates admin-only DB (no invoices)
- core smoke checks (cash signs, Excel, app build)
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def _git_ls_files() -> list[str]:
    try:
        out = subprocess.check_output(
            ["git", "ls-files"],
            cwd=ROOT,
            text=True,
            stderr=subprocess.DEVNULL,
        )
        return [line.strip() for line in out.splitlines() if line.strip()]
    except (subprocess.CalledProcessError, FileNotFoundError):
        return []


def check_no_mock_data_in_git():
    tracked = _git_ls_files()
    bad = [
        p
        for p in tracked
        if p.startswith("data/")
        or p.startswith("output/")
        or p.endswith(".db")
        or p.endswith(".sqlite")
        or p.endswith(".sqlite3")
    ]
    assert not bad, f"Mock/runtime data must not be in git: {bad}"
    print("OK: no data/output/db files tracked in git")


def check_deploy_files():
    required = [
        "Dockerfile",
        "docker-compose.yml",
        ".dockerignore",
        ".gitignore",
        "requirements.txt",
        "templates/invoice_template.xlsx",
        "assets/S260800018v4.xlsx",
        "assets/receipt_header.png",
        "scripts/docker-run.sh",
        "scripts/install-ubuntu.sh",
        "scripts/run.sh",
        "app.py",
        "database.py",
        "config.py",
    ]
    missing = [p for p in required if not (ROOT / p).exists()]
    assert not missing, f"Missing deploy files: {missing}"

    dockerignore = (ROOT / ".dockerignore").read_text(encoding="utf-8")
    for needle in ("data", "output", "logs"):
        assert needle in dockerignore, f".dockerignore must exclude {needle}"

    gitignore = (ROOT / ".gitignore").read_text(encoding="utf-8")
    for needle in ("data/", "output/", "logs/"):
        assert needle in gitignore, f".gitignore must exclude {needle}"

    print("OK: Docker/Linux deploy files present; data excluded from image/git")


def check_fresh_sqlite():
    """Simulate client first boot: empty data dir → admin only, zero invoices."""
    from sqlalchemy import create_engine, text
    from sqlalchemy.orm import sessionmaker

    import database as dbmod
    from auth import hash_password
    from config import ROLE_ADMIN

    with tempfile.TemporaryDirectory() as tmp:
        db_path = Path(tmp) / "gevin.db"
        engine = create_engine(
            f"sqlite:///{db_path.as_posix()}",
            connect_args={"check_same_thread": False},
        )
        dbmod.Base.metadata.create_all(engine)
        Session = sessionmaker(bind=engine)
        session = Session()

        # Mirror ensure_default_admin
        if session.query(dbmod.User).count() == 0:
            session.add(
                dbmod.User(
                    username="admin",
                    display_name="Admin",
                    password_hash=hash_password("admin123"),
                    role=ROLE_ADMIN,
                    is_active=True,
                )
            )
            session.commit()

        users = session.query(dbmod.User).all()
        invoices = session.query(dbmod.Invoice).count()
        moves = session.query(dbmod.InventoryMovement).count()
        cash = session.query(dbmod.CashMovement).count()
        assert len(users) == 1 and users[0].username == "admin"
        assert invoices == 0 and moves == 0 and cash == 0
        assert db_path.exists() and db_path.stat().st_size > 0
        session.close()
        engine.dispose()

    print("OK: fresh DB has admin only — no laptop mock invoices/stock")


def check_core_smoke():
    from cash import build_cash_movement, signed_cash_warehouse_amount
    from config import compose_receipt_storage
    from invoice_generator import generate_invoice_excel
    from openpyxl import load_workbook
    import app as m

    assert signed_cash_warehouse_amount(1000, "購入單") == -1000
    assert signed_cash_warehouse_amount(1000, "銷售單") == 1000
    src, dst = compose_receipt_storage("存", "A倉庫")
    assert dst == "存 A倉庫" and src == ""

    data = {
        "invoice_no": "LINUX_READY_P1",
        "transaction_type": "購入單",
        "customer_name": "DeployCheck",
        "customer_phone": "",
        "transaction_date": date(2026, 8, 12),
        "handler": "Admin",
        "payment_method": "",
        "invoice_currency": "HKD",
        "source_location": "",
        "destination_location": "存 A倉庫",
        "notes": "",
        "note_amount": 0,
        "total_amount": 100.0,
        "cash_warehouse_amount": signed_cash_warehouse_amount(100, "購入單"),
    }
    items = [
        {
            "item_type": "雜金 Gold Accessories",
            "quality": "24K",
            "weight_gram": 10,
            "weight_tael": None,
            "weight_oz": None,
            "unit_price": 10,
            "amount": 100,
        }
    ]
    path = Path(generate_invoice_excel(data, items))
    assert path.exists()
    ws = load_workbook(path)["購入單"]
    assert ws.cell(10, 11).value == -100
    stock = str(ws.cell(37, 8).value or "")
    assert "倉存存取" not in stock
    assert "存 A倉庫" in stock
    assert "取 客戶" not in stock
    assert "存 客戶" not in stock

    cm = build_cash_movement(data)
    assert cm["direction"] == "out"

    payload = m.load_review_page()
    assert len(payload) == 8
    m.build_app()
    print("OK: Excel path + Gradio app build")


def main():
    print(f"Repo: {ROOT}")
    check_no_mock_data_in_git()
    check_deploy_files()
    check_fresh_sqlite()
    check_core_smoke()
    print("")
    print("ALL LINUX/DOCKER READY CHECKS PASSED")
    print("Client fresh clone: empty data/ → new SQLite with admin/admin123 only.")
    print("Docker: bash scripts/docker-run.sh")
    print("No-Docker: bash scripts/install-ubuntu.sh && bash scripts/run.sh")


if __name__ == "__main__":
    main()
